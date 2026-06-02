"""OpenMIIR Excel Metadata Parser v3.9.4."""

import json
import os
import sys
from datetime import datetime, timezone

BASE = os.path.dirname(os.path.abspath(__file__))
CANDIDATES_DIR = os.path.join(BASE, "..", "..", "data", "external", "openmiir", "meta", "github_candidates")
EXPORTS = os.path.join(BASE, "..", "..", "data", "exports")

SEMANTIC_HEADERS = [
    "stimulus", "stimuli", "cue", "condition", "perception", "imagery",
    "listen", "imagine", "event", "marker", "trigger", "code", "id", "label",
    "block", "trial", "onset", "duration", "filename", "song", "fragment",
    "genre", "tempo", "meter", "lyrics", "name", "title", "description",
    "type", "category", "group", "set", "version", "perceive",
]


def _find_xlsx_files():
    files = []
    if not os.path.isdir(CANDIDATES_DIR):
        return files
    for fn in os.listdir(CANDIDATES_DIR):
        if any(fn.lower().endswith(e) for e in (".xlsx", ".xls", ".ods")):
            files.append(os.path.join(CANDIDATES_DIR, fn))
    return sorted(files)


def _has_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _has_pandas():
    try:
        import pandas  # noqa: F401
        return True
    except ImportError:
        return False


def parse_xlsx_openpyxl(path, original_name):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets_info = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = []
        row_count = 0
        sample_rows = []
        semantic_columns = []
        numeric_code_columns = []
        possible_mapping_rows = []

        for i, row in enumerate(ws.iter_rows(max_row=200, values_only=True)):
            if i == 0:
                headers = [str(c or "") for c in row]
                for j, h in enumerate(headers):
                    h_lower = h.lower().strip()
                    if any(kw in h_lower for kw in SEMANTIC_HEADERS):
                        semantic_columns.append({"index": j, "header": h, "match": h_lower})
                continue
            row_count += 1
            row_dict = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    row_dict[headers[j]] = cell

            if row_count <= 5:
                sample_rows.append({k: str(v) for k, v in list(row_dict.items())[:10]})

            has_code = False
            has_label = False
            for v in row_dict.values():
                if isinstance(v, (int, float)) and v > 0:
                    try:
                        code = int(v)
                        if code in [11, 12, 13, 14, 21, 22, 23, 24,
                                     31, 32, 33, 34, 41, 42, 43, 44,
                                     1000, 1111, 2000, 2001]:
                            has_code = True
                            if code not in numeric_code_columns:
                                numeric_code_columns.append(code)
                    except (ValueError, TypeError):
                        pass
                if isinstance(v, str) and any(
                    kw in v.lower() for kw in ["perception", "imagery", "listen", "imagine",
                                                "cue", "condition", "perceive"]
                ):
                    has_label = True

            if has_code and has_label:
                possible_mapping_rows.append({k: str(v) for k, v in list(row_dict.items())[:15]})

        sheets_info.append({
            "sheet": sheet_name,
            "headers": headers[:30],
            "row_count": row_count,
            "semantic_columns": semantic_columns,
            "sample_rows": sample_rows,
            "numeric_code_columns": numeric_code_columns,
            "possible_mapping_rows": possible_mapping_rows,
        })

    wb.close()
    return sheets_info


def parse_xlsx_pandas(path, original_name):
    import pandas as pd

    xl = pd.ExcelFile(path)
    sheets_info = []
    for sheet_name in xl.sheet_names:
        try:
            df = pd.read_excel(path, sheet_name=sheet_name, nrows=200)
        except Exception:
            continue
        headers = list(df.columns)
        row_count = len(df)
        sample_rows = [dict(df.iloc[i][:10]) for i in range(min(5, row_count))]
        for r in sample_rows:
            for k, v in r.items():
                if hasattr(v, 'item'):
                    r[k] = v.item()
                else:
                    r[k] = str(v)

        semantic_columns = []
        for j, h in enumerate(headers):
            h_lower = str(h).lower().strip()
            if any(kw in h_lower for kw in SEMANTIC_HEADERS):
                semantic_columns.append({"index": j, "header": str(h), "match": h_lower})

        numeric_code_columns = []
        possible_mapping_rows = []
        for _, row in df.iterrows():
            has_code = False
            has_label = False
            for v in row.values:
                if isinstance(v, (int, float)) and v > 0:
                    try:
                        code = int(v)
                        if code in [11, 12, 13, 14, 21, 22, 23, 24,
                                     31, 32, 33, 34, 41, 42, 43, 44,
                                     1000, 1111, 2000, 2001]:
                            has_code = True
                            if code not in numeric_code_columns:
                                numeric_code_columns.append(code)
                    except (ValueError, TypeError):
                        pass
                if isinstance(v, str) and any(
                    kw in v.lower() for kw in ["perception", "imagery", "listen", "imagine"]
                ):
                    has_label = True
            if has_code and has_label:
                possible_mapping_rows.append({str(k): str(v) for k, v in dict(row).items()})

        sheets_info.append({
            "sheet": sheet_name, "headers": headers[:30], "row_count": row_count,
            "semantic_columns": semantic_columns, "sample_rows": sample_rows,
            "numeric_code_columns": numeric_code_columns,
            "possible_mapping_rows": possible_mapping_rows,
        })
    return sheets_info


def main():
    os.makedirs(EXPORTS, exist_ok=True)
    xlsx_paths = _find_xlsx_files()
    workbooks = []
    files_parsed = 0

    use_openpyxl = _has_openpyxl()
    use_pandas = _has_pandas()

    if not use_openpyxl and not use_pandas:
        report = {
            "tool": "openmiir_excel_metadata_parser_v3.9.4",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "files_parsed": 0,
            "error": "No Excel parser available (openpyxl or pandas required)",
            "workbooks": [],
            "explicit_code_mappings": [],
            "stimulus_metadata": [],
            "condition_candidates": [],
            "evidence_level": "none",
        }
        print("No Excel parser available. Install openpyxl or pandas.", file=sys.stderr)
    else:
        parser_name = "openpyxl" if use_openpyxl else "pandas"
        for path in xlsx_paths:
            original = (os.path.basename(path).replace("_", "/", 1)
                       if "_" in os.path.basename(path)
                       else os.path.basename(path))
            try:
                if use_openpyxl:
                    sheets = parse_xlsx_openpyxl(path, original)
                else:
                    sheets = parse_xlsx_pandas(path, original)
                workbooks.append({"file": original, "sheets": sheets})
                files_parsed += 1
            except Exception as e:
                workbooks.append({"file": original, "error": str(e)})

        report = {
            "tool": "openmiir_excel_metadata_parser_v3.9.4",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "parser_used": parser_name,
            "files_parsed": files_parsed,
            "workbooks": workbooks,
            "explicit_code_mappings": [],
            "stimulus_metadata": [],
            "condition_candidates": [],
            "evidence_level": "none",
        }

        for wb in workbooks:
            for sh in wb.get("sheets", []):
                for row in sh.get("possible_mapping_rows", []):
                    report["explicit_code_mappings"].append({
                        "source": wb["file"], "sheet": sh["sheet"], "row": row,
                    })
                if sh.get("semantic_columns"):
                    report["stimulus_metadata"].append({
                        "source": wb["file"], "sheet": sh["sheet"],
                        "columns": [c["header"] for c in sh["semantic_columns"]],
                    })
        if report["explicit_code_mappings"]:
            report["evidence_level"] = "weak_hypothesis"

    json_path = os.path.join(EXPORTS, "openmiir_excel_metadata_index.json")
    with open(json_path, "w") as f:
        json.dump(report, f, indent=2, default=str)

    md_lines = [
        "# OpenMIIR Excel Metadata Index",
        f"**Files parsed**: {files_parsed}",
        f"**Parser**: {report.get('parser_used', 'none')}",
        f"**Evidence level**: {report['evidence_level']}",
    ]
    md_path = os.path.join(EXPORTS, "openmiir_excel_metadata_index.md")
    with open(md_path, "w") as f:
        f.write("\n".join(md_lines))

    print(f"Excel parser: files={files_parsed} evidence={report['evidence_level']}", file=sys.stderr)
    if report["explicit_code_mappings"]:
        print(f"  Code mappings found: {len(report['explicit_code_mappings'])}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
